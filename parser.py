"""
Parsers for Baer Shoes export documents:
 - Crystal-Reports style invoice (.xls / .xlsx)
 - ICEGATE Shipping Bill PDF ("LET EXPORT COPY")

Both return plain dicts / lists so they are easy to unit-test and to
feed into the Excel-writer module.
"""
import re
import subprocess
import tempfile
import os
import openpyxl
import pdfplumber

INVOICE_NO_RE = re.compile(r'([A-Z]{2,8}\d+-\d{5,6})\s*DT\s*(\d{2}-\d{2}-\d{4})')
SB_HEADER_RE  = re.compile(r'([A-Z]{2}[A-Z0-9]{3,4})\s+(\d{6,8})\s+(\d{2}-[A-Z]{3}-\d{2})')


# --------------------------------------------------------------------------
# 1. INVOICE PARSER
# --------------------------------------------------------------------------
def _to_xlsx_if_needed(path: str) -> str:
    """Legacy .xls (Crystal Reports / BIFF) -> .xlsx via headless LibreOffice,
    so we can read it with openpyxl without needing xlrd."""
    if path.lower().endswith(".xlsx"):
        return path
    out_dir = tempfile.mkdtemp()
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out_dir, path],
        check=True, capture_output=True,
    )
    base = os.path.splitext(os.path.basename(path))[0]
    return os.path.join(out_dir, base + ".xlsx")


def parse_invoice(path: str) -> dict:
    xlsx_path = _to_xlsx_if_needed(path)
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    invoice_no = invoice_date = exchange_rate = None
    buyer = country = None
    line_items = []
    current_desc = None

    for row in ws.iter_rows():
        for cell in row:
            v = cell.value
            if v is None:
                continue
            s = str(v)

            # Invoice number + date, e.g. " BG047-202627 DT 05-09-2026"
            if invoice_no is None:
                m = INVOICE_NO_RE.search(s)
                if m:
                    invoice_no, invoice_date = m.group(1), m.group(2)

            # Exchange rate label -> numeric value in same row
            if s.strip().startswith("Exchange Rate / Per Euro Rs"):
                nums = [c.value for c in row if isinstance(c.value, (int, float))]
                if nums:
                    exchange_rate = nums[0]

            # Buyer / consignee block
            if s.strip().upper().startswith("CONSIGNEE"):
                lines = s.split("\n")
                if len(lines) > 1:
                    buyer = lines[1].strip()

            if s.strip().upper() == "GERMANY" and country is None:
                country = "GERMANY"

            # Article description row, e.g. "Leather - Full Shoes for Adults (Gents)"
            if "for Adults" in s or "FOR ADULTS" in s.upper():
                current_desc = s.strip()

            # Grade row -> qty / rate / value are the numeric cells on this row
            if s.strip().lower().startswith("grade") and current_desc:
                nums = sorted(
                    ((c.column, c.value) for c in row if isinstance(c.value, (int, float))),
                    key=lambda t: t[0],
                )
                if len(nums) >= 3:
                    qty, rate, value = nums[0][1], nums[1][1], nums[2][1]
                    grade = "II" if "2" in s else "I"
                    line_items.append({
                        "description": current_desc,
                        "grade": grade,
                        "qty": qty,
                        "rate_eur": rate,
                        "value_eur": round(value, 2),
                    })

    return {
        "invoice_no": invoice_no,
        "invoice_date": invoice_date,      # string DD-MM-YYYY, converted later
        "exchange_rate": exchange_rate,
        "buyer": buyer or "BAR GMBH",
        "country": country or "GERMANY",
        "line_items": line_items,
    }


# --------------------------------------------------------------------------
# 2. SHIPPING BILL (PDF) PARSER
# --------------------------------------------------------------------------
def parse_shipping_bill(path: str) -> dict:
    with pdfplumber.open(path) as pdf:
        text = pdf.pages[0].extract_text()

    sb_no = sb_date = port_code = None
    m = SB_HEADER_RE.search(text)
    if m:
        port_code, sb_no, sb_date = m.group(1), int(m.group(2)), m.group(3)

    inv_match = re.search(r'\d+\s+([A-Z]{2,8}\d+-\d{5,6})\s+([\d.]+)\s+EUR', text)
    invoice_no = inv_match.group(1) if inv_match else None
    invoice_amt = float(inv_match.group(2)) if inv_match else None

    return {
        "sb_no": sb_no,
        "sb_date": sb_date,          # string DD-MMM-YY
        "port_code": port_code,
        "invoice_no": invoice_no,    # used to match SB back to its invoice
        "invoice_amount_eur": invoice_amt,
    }
