"""
Appends parsed invoice + shipping-bill data into the 'Export Sales' sheet
of the master checklist workbook, in EXACTLY the same layout the company
already uses (verified against the last real entries in the file):

  SLNO | DATE | INVOICE NO | SB NO | SB DATE | COUNTRY | Customer |
  CATEGORY | PAIRS | VALUE(EUR) | EX RATE | VALUE(INR) | .. | Month

- VALUE(INR) is written as a live formula '=K{r}*J{r}' (matches current
  practice), so it recalculates itself -- never a hard-coded number.
- Every invoice group ends with a subtotal row using SUM()/reference
  formulas, exactly like every existing group in the sheet.
- Every other sheet (Export Sales Summary, etc.) is left completely
  untouched: its SUMIFS formulas already point at the 'Export Sales'
  range and will simply pick up the new rows the next time Excel
  recalculates (F9) or the file is opened.
"""
import datetime as dt
from copy import copy as copy_style
import openpyxl
from categorize import map_category

SHEET = "Export Sales"
NUM_COLS = 17  # columns A..Q carry the row's font/fill/border styling


def _apply_style(dst_cell, src_cell):
    """Copy the full visual style (font, fill, border, alignment,
    number format) from an existing template cell onto a new one, so
    appended rows are visually identical to the rest of the sheet --
    including themed fills like the coloured subtotal-row background,
    which a plain RGB copy would miss."""
    dst_cell.font = copy_style(src_cell.font)
    dst_cell.fill = copy_style(src_cell.fill)
    dst_cell.border = copy_style(src_cell.border)
    dst_cell.alignment = copy_style(src_cell.alignment)
    dst_cell.protection = copy_style(src_cell.protection)
    dst_cell.number_format = src_cell.number_format


def _parse_ddmmyyyy(s: str) -> dt.datetime:
    return dt.datetime.strptime(s, "%d-%m-%Y")


def _parse_ddmonyy(s: str) -> dt.datetime:
    return dt.datetime.strptime(s, "%d-%b-%y")


def append_invoice(wb: openpyxl.Workbook, invoice: dict, sb: dict) -> list:
    """Returns the list of category rows written, for a quick on-screen review."""
    ws = wb[SHEET]

    last_row = ws.max_row
    # The existing last row is always a subtotal row; the row just above
    # it is always a normal data row. Use both as style templates so new
    # rows exactly match (bold + themed fill for subtotals, plain for data).
    template_subtotal_row = last_row
    template_data_row = last_row - 1

    # find current max SLNO (first column) scanning upward from the bottom
    max_slno = 0
    for r in range(last_row, 4, -1):
        v = ws.cell(row=r, column=1).value
        if v not in (None, ""):
            try:
                max_slno = int(v)
            except (TypeError, ValueError):
                pass
            break
    next_slno = max_slno + 1

    invoice_date = _parse_ddmmyyyy(invoice["invoice_date"])
    sb_date = _parse_ddmonyy(sb["sb_date"]) if sb.get("sb_date") else None
    month = dt.datetime(invoice_date.year, invoice_date.month, 1)

    start_row = last_row + 1
    r = start_row
    written = []

    for item in invoice["line_items"]:
        category = map_category(item["description"], item["grade"])
        ws.cell(row=r, column=1, value=next_slno)                 # A SLNO
        ws.cell(row=r, column=2, value=invoice_date)               # B DATE
        ws.cell(row=r, column=3, value=invoice["invoice_no"])      # C INVOICE NO
        ws.cell(row=r, column=4, value=sb.get("sb_no"))            # D SB NO
        ws.cell(row=r, column=5, value=sb_date)                    # E SB DATE
        ws.cell(row=r, column=6, value=invoice["country"])         # F COUNTRY
        ws.cell(row=r, column=7, value=invoice["buyer"])           # G Customer
        ws.cell(row=r, column=8, value=category)                   # H CATEGORY
        ws.cell(row=r, column=9, value=item["qty"])                 # I PAIRS
        ws.cell(row=r, column=10, value=item["value_eur"])          # J VALUE(EUR)
        ws.cell(row=r, column=11, value=invoice["exchange_rate"])   # K EX RATE
        ws.cell(row=r, column=12, value=f"=K{r}*J{r}")              # L VALUE(INR)
        ws.cell(row=r, column=17, value=month)                      # Q Month

        for c in range(1, NUM_COLS + 1):
            _apply_style(ws.cell(row=r, column=c), ws.cell(row=template_data_row, column=c))

        written.append((r, category, item["qty"], item["value_eur"]))
        r += 1

    end_row = r - 1
    ws.cell(row=r, column=9, value=f"=SUM(I{start_row}:I{end_row})")
    ws.cell(row=r, column=10, value=f"=SUM(J{start_row}:J{end_row})")
    ws.cell(row=r, column=12, value=f"=SUM(L{start_row}:L{end_row})")
    ws.cell(row=r, column=17, value=f"=Q{end_row}")

    for c in range(1, NUM_COLS + 1):
        _apply_style(ws.cell(row=r, column=c), ws.cell(row=template_subtotal_row, column=c))

    # match row height to the template rows too, if they were customized
    ws.row_dimensions[r].height = ws.row_dimensions[template_subtotal_row].height
    for rr in range(start_row, end_row + 1):
        ws.row_dimensions[rr].height = ws.row_dimensions[template_data_row].height

    return written
