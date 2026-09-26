"""
Auto-extends 'Export Sales Summary' with a new month's 5-column formula
block, by cloning the most recent existing block.

Each month occupies 5 columns: Pairs | Value(EUR) | Ex.Rate | Value(INR) |
Per Pair. The 'Value(INR)' column also carries the month date in row 3,
and every SUMIFS formula in the block references that cell (as an
absolute reference, e.g. $CG$3) as its month filter.

Cloning strategy: copy every formula in the template block cell-for-cell,
and do a targeted column-letter substitution -- replacing only the
block's own 5 column letters (e.g. CD,CE,CF,CG,CH -> CI,CJ,CK,CL,CM) --
so cross-sheet references to 'Export Sales' (I:I, J:J, L:L, H:H, Q:Q,
and the category-label column A) are left completely untouched. This
preserves the sheet's existing per-row quirks exactly (some rows have
no 'Per Pair' formula, row 25's SUM range differs from row 21's, etc.)
because we copy the template's actual formulas rather than assuming a
uniform pattern.
"""
import re
import datetime as dt
from copy import copy as copy_style
from openpyxl.utils import get_column_letter, column_index_from_string

SHEET = "Export Sales Summary"
BLOCK_WIDTH = 5
MONTH_ROW = 3
FIRST_ROW = 3
LAST_ROW = 36


def _find_last_block_anchor_col(ws) -> int:
    """The anchor column of a block is its 4th column (Value INR), which
    carries the month date in row MONTH_ROW. Find the right-most one."""
    last_col = None
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=MONTH_ROW, column=c).value
        if isinstance(v, dt.datetime):
            last_col = c
    if last_col is None:
        raise ValueError("No existing month block found in Export Sales Summary")
    return last_col


def _add_month(month_dt: dt.datetime) -> dt.datetime:
    y, m = month_dt.year, month_dt.month
    return dt.datetime(y + (1 if m == 12 else 0), 1 if m == 12 else m + 1, 1)


def _clone_block(ws, old_anchor_col: int) -> int:
    old_start = old_anchor_col - 3
    new_start = old_start + BLOCK_WIDTH
    old_cols = [get_column_letter(old_start + i) for i in range(BLOCK_WIDTH)]
    new_cols = [get_column_letter(new_start + i) for i in range(BLOCK_WIDTH)]

    # Replace whole column-letter tokens only (not part of a longer ref,
    # not preceded by another letter), for the 5 block columns only.
    col_pattern = re.compile(
        r'(?<![A-Za-z])(' + '|'.join(sorted(old_cols, key=len, reverse=True)) + r')(?=\$?\d)'
    )
    col_map = dict(zip(old_cols, new_cols))

    def shift_formula(formula: str) -> str:
        return col_pattern.sub(lambda m: col_map[m.group(1)], formula)

    old_month = ws.cell(row=MONTH_ROW, column=old_anchor_col).value
    new_month = _add_month(old_month)

    for row in range(FIRST_ROW, LAST_ROW + 1):
        for i in range(BLOCK_WIDTH):
            old_c = old_start + i
            new_c = new_start + i
            src = ws.cell(row=row, column=old_c)
            dst = ws.cell(row=row, column=new_c)

            if row == MONTH_ROW and (old_c) == old_anchor_col:
                dst.value = new_month
            elif isinstance(src.value, str) and src.value.startswith("="):
                dst.value = shift_formula(src.value)
            else:
                dst.value = src.value  # literal headers, '-' placeholders, etc.

            # visual style: font, fill, border, number format
            dst.font = copy_style(src.font)
            dst.fill = copy_style(src.fill)
            dst.border = copy_style(src.border)
            dst.alignment = copy_style(src.alignment)
            dst.number_format = src.number_format

        ws.row_dimensions[row].height = ws.row_dimensions[row].height

    for i in range(BLOCK_WIDTH):
        old_letter = get_column_letter(old_start + i)
        new_letter = get_column_letter(new_start + i)
        if old_letter in ws.column_dimensions:
            ws.column_dimensions[new_letter].width = ws.column_dimensions[old_letter].width

    return new_start + 3  # new anchor column index


def ensure_month_block(wb, target_month: dt.datetime) -> None:
    """Make sure Export Sales Summary has a formula block for target_month
    (first-of-month datetime). Clones forward one month at a time from
    whatever the latest existing block is, so gaps of more than one
    missing month are filled correctly in sequence."""
    ws = wb[SHEET]
    target_month = dt.datetime(target_month.year, target_month.month, 1)

    anchor_col = _find_last_block_anchor_col(ws)
    current_month = ws.cell(row=MONTH_ROW, column=anchor_col).value

    guard = 0
    while current_month < target_month:
        anchor_col = _clone_block(ws, anchor_col)
        current_month = ws.cell(row=MONTH_ROW, column=anchor_col).value
        guard += 1
        if guard > 60:  # sanity stop: 5 years of monthly blocks
            raise RuntimeError("Too many months to backfill -- check the target date")
