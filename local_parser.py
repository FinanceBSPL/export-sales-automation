"""
Parses Baer Shoes' domestic ("TAX-Invoice") local sales invoices, and the
small LUT/courier export invoices that -- per company convention -- are
still logged in the 'Local Sales working' sheet (e.g. Fontana via DHL).

Both return: invoice_no, invoice_date, customer, goods_description,
pairs, value_inr, value_eur (None for pure-domestic invoices),
exchange_rate (None for pure-domestic invoices).
"""
import re
import subprocess
import tempfile
import os
import openpyxl

INV_NO_DATE_RE = re.compile(r'([A-Z0-9]+-\d{3,5})\s*/\s*(\d{2}-\d{2}-\d{4})')
INV_NO_DATE_RE2 = re.compile(r'([A-Z0-9]+-\d{3,5})\s*dtd\s*(\d{2}\.\d{2}\.\d{4})', re.IGNORECASE)


def _to_xlsx_if_needed(path: str) -> str:
    if path.lower().endswith(".xlsx"):
        return path
    out_dir = tempfile.mkdtemp()
    subprocess.run(
        ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", out_dir, path],
        check=True, capture_output=True,
    )
    base = os.path.splitext(os.path.basename(path))[0]
    return os.path.join(out_dir, base + ".xlsx")


def _find_first_sheet_with_content(wb):
    for sn in wb.sheetnames:
        ws = wb[sn]
        if ws.max_row and ws.max_row > 5:
            return ws
    return wb[wb.sheetnames[0]]


def parse_domestic_invoice(path: str) -> dict:
    """Baer's 'TAX-Invoice' template (IGST/CGST+SGST, Indian buyer)."""
    xlsx_path = _to_xlsx_if_needed(path)
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = _find_first_sheet_with_content(wb)

    invoice_no = invoice_date = customer = goods_description = None
    pairs = value_inr = None

    rows = list(ws.iter_rows())
    for i, row in enumerate(rows):
        for cell in row:
            v = cell.value
            if v is None:
                continue
            s = str(v)

            if invoice_no is None:
                m = INV_NO_DATE_RE.search(s)
                if m:
                    invoice_no, invoice_date = m.group(1), m.group(2)

            if s.strip().lower().startswith("bill to") and customer is None:
                # customer name is the next non-empty cell below, same column
                for j in range(i + 1, min(i + 3, len(rows))):
                    v2 = rows[j][cell.column - 1].value
                    if v2:
                        customer = str(v2).replace("M/s.", "").strip()
                        break

            if s.strip().lower().startswith("description of goods") and goods_description is None:
                for j in range(i + 1, min(i + 3, len(rows))):
                    v2 = rows[j][cell.column - 1].value
                    if v2:
                        goods_description = str(v2).strip()
                        break

            # the pre-GST 'Total' line: exact match, not 'Total Amount'
            if s.strip() == "Total":
                nums = [c.value for c in row if isinstance(c.value, (int, float))]
                if len(nums) >= 2:
                    pairs, value_inr = nums[0], nums[-1]

    return {
        "invoice_no": invoice_no,
        "invoice_date": invoice_date,       # DD-MM-YYYY
        "customer": customer,
        "goods_description": goods_description,
        "pairs": pairs,
        "value_inr": value_inr,
        "value_eur": None,
        "exchange_rate": None,
    }


def parse_lut_courier_invoice(path: str) -> dict:
    """Small LUT/zero-rated export invoices (e.g. DHL courier samples) that
    the company still logs inside 'Local Sales working' under a specific
    customer's table, rather than through the main Export Sales flow."""
    xlsx_path = _to_xlsx_if_needed(path)
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = _find_first_sheet_with_content(wb)

    invoice_no = invoice_date = customer = goods_description = None
    pairs = value_eur = value_inr = exchange_rate = None

    rows = list(ws.iter_rows())
    for i, row in enumerate(rows):
        for cell in row:
            v = cell.value
            if v is None:
                continue
            s = str(v)

            if invoice_no is None:
                m = INV_NO_DATE_RE2.search(s)
                if m:
                    invoice_no, invoice_date = m.group(1), m.group(2).replace(".", "-")

            if s.strip().lower() == "consignee" and customer is None:
                for j in range(i + 1, min(i + 3, len(rows))):
                    v2 = rows[j][cell.column - 1].value
                    if v2:
                        customer = str(v2).strip()
                        break

            if s.strip().lower().startswith("description of goods") and goods_description is None:
                for j in range(i + 1, min(i + 3, len(rows))):
                    v2 = rows[j][cell.column - 1].value
                    if v2:
                        goods_description = str(v2).strip()
                        break

            if s.strip().lower().startswith("exchange rate"):
                nums = [c.value for c in row if isinstance(c.value, (int, float))]
                if nums:
                    exchange_rate = nums[0]

            if "fob. rs" in s.strip().lower():
                nums = [c.value for c in row if isinstance(c.value, (int, float))]
                if nums:
                    value_inr = nums[0]

            if s.strip().lower().startswith("amount chargeable"):
                nums = [c.value for c in row if isinstance(c.value, (int, float))]
                if len(nums) >= 2:
                    pairs, value_eur = nums[0], nums[-1]

    return {
        "invoice_no": invoice_no,
        "invoice_date": invoice_date,
        "customer": customer,
        "goods_description": goods_description,
        "pairs": pairs,
        "value_inr": value_inr,
        "value_eur": value_eur,
        "exchange_rate": exchange_rate,
    }
